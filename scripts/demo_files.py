"""Demo documents, built in memory: a BOM (two versions), a cable spec and a connector spec."""

import io

import docx
import openpyxl
import pymupdf

HEADER = ["Item", "Material No", "Description", "Supplier", "Qty"]

_ROWS = [
    ("J1", "MN-1001", "Power input connector, 2-pin", "Harlow Components", 1),
    ("J2", "MN-1002", "Sensor connector, 4-pin", "Harlow Components", 2),
    ("J3", "MN-1003", "Debug header connector, 6-pin", "Harlow Components", 1),
    ("J4", "MN-1004", "Board-to-cable connector, 10-pin", "Northgate Connectors", 1),
    ("C1", "MN-2001", "Cable assembly, 300 mm, 10-conductor", "Vale Cable Works", 1),
    ("U1", "MN-3001", "Microcontroller module", "Brightline Semiconductor", 1),
    ("R1", "MN-4001", "Resistor, 10 kOhm, 0603", "Delta Passives", 12),
    ("H1", "MN-5001", "Enclosure, injection molded", "Lakeside Plastics", 1),
]

NEW_J4_SUPPLIER = "Ridgeview Connectors"


def _save(wb: openpyxl.Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def bom_xlsx(version: int = 1) -> bytes:
    """Version 2 swaps the J4 connector supplier and adds a Changes sheet."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOM"
    ws.append(HEADER)
    for item, material, description, supplier, qty in _ROWS:
        if version >= 2 and item == "J4":
            supplier = NEW_J4_SUPPLIER
        ws.append([item, material, description, supplier, qty])
    if version >= 2:
        changes = wb.create_sheet("Changes")
        changes.append(["Item", "Change"])
        changes.append(
            [
                "J4",
                "Connector J4 supplier changed from Northgate Connectors to "
                f"{NEW_J4_SUPPLIER}: the original part is end-of-life.",
            ]
        )
    return _save(wb)


def cable_spec_docx() -> bytes:
    d = docx.Document()
    d.add_heading("Cable Assembly Spec", level=1)
    d.add_paragraph(
        "Scope: 300 mm, 10-conductor cable assembly joining the main board to the sensor module."
    )
    d.add_paragraph("The cable shall be terminated with the board-to-cable connector (J4).")
    d.add_paragraph("Conductors: 28 AWG stranded copper, PVC insulation, rated 300 V.")
    d.add_paragraph("Each cable assembly shall pass a continuity test before shipment.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def connector_spec_pdf() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    lines = [
        "Supplier Spec - Connector",
        "",
        "Applies to: board-to-cable connector J4 (MN-1004), 10-pin.",
        "Supplier: Northgate Connectors.",
        "Contact plating: 30 microinch gold. Rated current: 1 A per contact.",
        "Incoming inspection: 5-piece sample per lot, pin gauge and visual.",
    ]
    for i, line in enumerate(lines):
        page.insert_text((72, 72 + 18 * i), line)
    data: bytes = doc.tobytes()
    doc.close()
    return data
