import io
import zipfile

import docx
import openpyxl

from services import extract
from services.extract import DOCX, PDF, XLSX, extract_text, sniff_type

pymupdf = extract.pymupdf


def _pdf(text: str = "hello pdf") -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


def _encrypted_pdf() -> bytes:
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "secret")
    data = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="x", owner_pw="x")
    doc.close()
    return data


def _docx() -> bytes:
    document = docx.Document()
    document.add_paragraph("Paragraph text")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Cell text"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Part"
    ws["B1"] = "J4"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_sniff_pdf():
    assert sniff_type(_pdf()) == PDF


def test_sniff_docx_vs_xlsx():
    assert sniff_type(_docx()) == DOCX
    assert sniff_type(_xlsx()) == XLSX


def test_sniff_rejects_zip_without_office_parts():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("a.txt", "x")
    assert sniff_type(buf.getvalue()) is None


def test_sniff_rejects_html():
    assert sniff_type(b"<html><body>hi</body></html>") is None


def test_extract_xlsx_cells_tab_separated():
    text = extract_text(XLSX, _xlsx())
    assert text is not None
    assert "Part\tJ4" in text


def test_extract_docx_includes_table_cells():
    text = extract_text(DOCX, _docx())
    assert text is not None
    assert "Paragraph text" in text
    assert "Cell text" in text


def test_extract_encrypted_pdf_returns_none():
    assert extract_text(PDF, _encrypted_pdf()) is None


def test_extract_truncates_at_cap(monkeypatch):
    monkeypatch.setattr(extract, "MAX_TEXT_CHARS", 10)
    text = extract_text(PDF, _pdf("abcdefghijklmnopqrstuvwxyz"))
    assert text is not None
    assert len(text) == 10
