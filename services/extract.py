import io
import warnings
import zipfile

import docx
import openpyxl

with warnings.catch_warnings():
    # pymupdf's SWIG bindings emit DeprecationWarnings at import time.
    warnings.simplefilter("ignore", DeprecationWarning)
    import pymupdf

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ACCEPTED = frozenset({PDF, DOCX, XLSX})

MAX_TEXT_CHARS = 2_097_152


def sniff_type(data: bytes) -> str | None:
    if data.startswith(b"%PDF-"):
        return PDF
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                names = set(z.namelist())
        except zipfile.BadZipFile:
            return None
        if "word/document.xml" in names:
            return DOCX
        if "xl/workbook.xml" in names:
            return XLSX
    return None


def _pdf_text(data: bytes) -> str:
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        if doc.needs_pass:
            raise ValueError("encrypted pdf")
        return "\n".join(page.get_text() for page in doc)


def _docx_text(data: bytes) -> str:
    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _xlsx_text(data: bytes) -> str:
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        rows: list[str] = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None and str(c) != ""]
                if cells:
                    rows.append("\t".join(cells))
        return "\n".join(rows)
    finally:
        wb.close()


_EXTRACTORS = {PDF: _pdf_text, DOCX: _docx_text, XLSX: _xlsx_text}


def extract_text(content_type: str, data: bytes) -> str | None:
    extractor = _EXTRACTORS.get(content_type)
    if extractor is None:
        return None
    try:
        text = extractor(data)
    except Exception:
        return None
    text = text[:MAX_TEXT_CHARS]
    if not text.strip():
        return None
    return text
